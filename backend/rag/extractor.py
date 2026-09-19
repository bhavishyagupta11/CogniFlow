"""
Unified Document Extraction Abstraction for CogniFlow
Supports:
1. PDF (.pdf) - PyMuPDF with password support & in-memory decryption
2. DOCX (.docx) - python-docx preserving paragraphs, headings, and tables in document order
3. TXT (.txt) - UTF-8 decoding with robust fallback handling
4. MD (.md) - Markdown text preserving syntax and header-based sectioning
"""

import io
import re
from pathlib import Path
from typing import List, Dict, Any, Tuple, Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
    from docx.text.paragraph import Paragraph
    from docx.table import Table
except ImportError:
    docx = None


class PDFPasswordRequiredError(Exception):
    """Raised when an encrypted PDF is uploaded without a password."""
    pass


class PDFIncorrectPasswordError(Exception):
    """Raised when the provided password fails to decrypt the encrypted PDF."""
    pass


class UnsupportedFormatError(Exception):
    """Raised when an uploaded file format is not supported."""
    pass


class ExtractionError(Exception):
    """Raised when text extraction fails."""
    pass


def sniff_document_type(file_bytes: bytes, filename: str, mime_type: Optional[str] = None) -> str:
    """
    Determines normalized format ('pdf', 'docx', 'txt', 'md') from extension, MIME, and magic bytes.
    Raises UnsupportedFormatError for unsupported formats like legacy .doc.
    """
    ext = Path(filename).suffix.lower()

    # Check for legacy binary .doc
    if ext == ".doc":
        raise UnsupportedFormatError(
            "Legacy binary .doc files are not supported. Please save or convert your document to modern .docx format before uploading."
        )

    # Magic byte checks
    if file_bytes.startswith(b"%PDF-"):
        return "pdf"

    if file_bytes.startswith(b"PK\x03\x04"):
        # Zip archive - check if it's a docx
        try:
            import zipfile
            with zipfile.ZipFile(io.BytesIO(file_bytes)) as zf:
                namelist = zf.namelist()
                if "word/document.xml" in namelist:
                    return "docx"
        except Exception:
            pass

    # Extension-based mapping
    if ext == ".pdf":
        return "pdf"
    if ext == ".docx":
        return "docx"
    if ext == ".md" or ext == ".markdown":
        return "md"
    if ext in [".txt", ".text", ".log", ".csv", ".json"]:
        return "txt"

    # MIME-based mapping
    if mime_type:
        mime = mime_type.lower()
        if "pdf" in mime:
            return "pdf"
        if "wordprocessingml" in mime or "officedocument" in mime:
            return "docx"
        if "markdown" in mime:
            return "md"
        if "text/" in mime:
            return "txt"

    raise UnsupportedFormatError(
        f"Unsupported file format '{ext or filename}'. Supported formats are: .pdf, .docx, .txt, .md."
    )


def extract_pdf_content(
    file_bytes: bytes,
    filename: str,
    password: Optional[str] = None
) -> Tuple[List[Dict[str, Any]], bytes]:
    """
    Extracts text page-by-page from a PDF file using PyMuPDF.
    Returns: (sections, unlocked_bytes)
    """
    if fitz is None:
        raise RuntimeError("PyMuPDF (fitz) is not installed.")

    doc = fitz.open(stream=file_bytes, filetype="pdf")
    try:
        was_encrypted = False
        if doc.needs_pass:
            if not password:
                raise PDFPasswordRequiredError(
                    "This PDF is password-protected. Please enter the password to unlock and process it."
                )
            auth_res = doc.authenticate(password)
            if auth_res == 0:
                raise PDFIncorrectPasswordError("Incorrect PDF password. Please try again.")
            was_encrypted = True

        sections = []
        for i, page in enumerate(doc):
            text = page.get_text("text") or ""
            # Reconnect numbered lists like "1. \nTitle" -> "1. Title"
            text = re.sub(r"(\b\d{1,3}\.)\s*\n\s*", r"\1 ", text)
            # Reconnect bullet lists like "• \nItem" -> "• Item"
            text = re.sub(r"([•\-\*])\s*\n\s*", r"\1 ", text)
            # Clean extra horizontal whitespace
            cleaned = re.sub(r"[ \t]+", " ", text).strip()
            if cleaned:
                # Detect first line as tentative section if heading-like
                first_line = cleaned.split("\n")[0].strip()
                sec_title = f"Page {i + 1}"
                if len(first_line) < 80 and not first_line.endswith("."):
                    sec_title = first_line

                sections.append({
                    "page_number": i + 1,
                    "pageNumber": i + 1,
                    "section": sec_title,
                    "text": cleaned,
                    "metadata": {
                        "page": i + 1,
                        "format": "pdf",
                        "character_count": len(cleaned)
                    }
                })

        unlocked_bytes = doc.tobytes() if was_encrypted else file_bytes
        return sections, unlocked_bytes
    finally:
        doc.close()


def extract_docx_content(
    file_bytes: bytes,
    filename: str
) -> Tuple[List[Dict[str, Any]], bytes]:
    """
    Extracts structured content from a DOCX file using python-docx.
    Traverses paragraphs, headings, and tables in strict document order.
    Returns: (sections, file_bytes)
    """
    if docx is None:
        raise RuntimeError("python-docx is not installed.")

    try:
        doc = docx.Document(io.BytesIO(file_bytes))
    except Exception as e:
        raise ExtractionError(f"Failed to read DOCX file: {e}")

    stem_name = Path(filename).stem
    sections: List[Dict[str, Any]] = []
    current_section_title = stem_name
    current_paragraphs: List[str] = []

    def flush_section():
        nonlocal current_paragraphs, current_section_title
        if not current_paragraphs:
            return
        combined_text = "\n\n".join(current_paragraphs).strip()
        if combined_text:
            sections.append({
                "page_number": None,
                "pageNumber": None,
                "section": current_section_title,
                "text": combined_text,
                "metadata": {
                    "format": "docx",
                    "section_title": current_section_title,
                    "character_count": len(combined_text)
                }
            })
        current_paragraphs = []

    # Iterate through elements in document body order
    for child in doc.element.body:
        if child.tag.endswith("p"):
            p = Paragraph(child, doc)
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style and p.style.name else ""
            is_heading = "heading" in style_name or style_name.startswith("title")

            # Text-based heading heuristic if standard heading style not used
            if not is_heading and len(text) < 70 and not text.endswith("."):
                if text.isupper() or re.match(r"^(?:chapter|section|\d+\.\d+|\d+\.)\s+", text, re.IGNORECASE):
                    is_heading = True

            if is_heading:
                flush_section()
                current_section_title = text
                current_paragraphs.append(f"## {text}")
            else:
                current_paragraphs.append(text)

        elif child.tag.endswith("tbl"):
            t = Table(child, doc)
            table_lines = []
            for row_idx, row in enumerate(t.rows):
                row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                if any(row_cells):
                    table_lines.append("| " + " | ".join(row_cells) + " |")
                    if row_idx == 0:
                        table_lines.append("| " + " | ".join(["---"] * len(row_cells)) + " |")
            if table_lines:
                current_paragraphs.append("\n".join(table_lines))

    flush_section()

    if not sections and len(doc.paragraphs) > 0:
        all_text = "\n\n".join(p.text.strip() for p in doc.paragraphs if p.text.strip())
        if all_text:
            sections.append({
                "page_number": None,
                "pageNumber": None,
                "section": stem_name,
                "text": all_text,
                "metadata": {"format": "docx", "section_title": stem_name, "character_count": len(all_text)}
            })

    return sections, file_bytes


def extract_txt_content(
    file_bytes: bytes,
    filename: str
) -> Tuple[List[Dict[str, Any]], bytes]:
    """
    Extracts text from a plain text file using UTF-8 with fallbacks.
    Partitions into logical sections based on double newlines and header patterns.
    Returns: (sections, file_bytes)
    """
    raw_text = None
    encodings_to_try = ["utf-8", "utf-8-sig", "latin-1", "cp1252"]
    for enc in encodings_to_try:
        try:
            raw_text = file_bytes.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue

    if raw_text is None:
        raw_text = file_bytes.decode("utf-8", errors="replace")

    stem_name = Path(filename).stem
    normalized = re.sub(r"\r\n?", "\n", raw_text)
    paragraphs = re.split(r"\n\s*\n", normalized)

    sections: List[Dict[str, Any]] = []
    current_section = stem_name
    current_paras: List[str] = []
    current_chars = 0

    def flush_txt_section():
        nonlocal current_paras, current_section, current_chars
        if not current_paras:
            return
        combined = "\n\n".join(current_paras).strip()
        if combined:
            sections.append({
                "page_number": None,
                "pageNumber": None,
                "section": current_section,
                "text": combined,
                "metadata": {
                    "format": "txt",
                    "section_title": current_section,
                    "character_count": len(combined)
                }
            })
        current_paras = []
        current_chars = 0

    for para in paragraphs:
        para_str = para.strip()
        if not para_str:
            continue

        first_line = para_str.split("\n")[0].strip()
        # Detect section header
        if len(first_line) <= 70 and (
            first_line.isupper()
            or re.match(r"^(?:unit|chapter|section|\d+\.\d+|\d+\.)\s+", first_line, re.IGNORECASE)
            or (first_line.endswith(":") and len(first_line) < 50)
        ):
            if current_chars >= 800:
                flush_txt_section()
            current_section = first_line.rstrip(":")

        current_paras.append(para_str)
        current_chars += len(para_str)
        if current_chars >= 1800:
            flush_txt_section()

    flush_txt_section()

    if not sections and raw_text.strip():
        sections.append({
            "page_number": None,
            "pageNumber": None,
            "section": stem_name,
            "text": raw_text.strip(),
            "metadata": {"format": "txt", "section_title": stem_name, "character_count": len(raw_text.strip())}
        })

    return sections, file_bytes


def extract_markdown_content(
    file_bytes: bytes,
    filename: str
) -> Tuple[List[Dict[str, Any]], bytes]:
    """
    Extracts text from a Markdown file, preserving formatting and splitting by markdown headers.
    Returns: (sections, file_bytes)
    """
    try:
        raw_text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        raw_text = file_bytes.decode("utf-8", errors="replace")

    stem_name = Path(filename).stem
    normalized = re.sub(r"\r\n?", "\n", raw_text)
    lines = normalized.split("\n")

    sections: List[Dict[str, Any]] = []
    current_section = stem_name
    current_lines: List[str] = []

    def flush_md_section():
        nonlocal current_lines, current_section
        if not current_lines:
            return
        combined = "\n".join(current_lines).strip()
        if combined:
            sections.append({
                "page_number": None,
                "pageNumber": None,
                "section": current_section,
                "text": combined,
                "metadata": {
                    "format": "md",
                    "section_title": current_section,
                    "character_count": len(combined)
                }
            })
        current_lines = []

    for line in lines:
        match = re.match(r"^(#{1,4})\s+(.+)$", line.strip())
        if match:
            # New markdown heading
            flush_md_section()
            current_section = match.group(2).strip()
            current_lines.append(line)
        else:
            current_lines.append(line)

    flush_md_section()

    if not sections and raw_text.strip():
        sections.append({
            "page_number": None,
            "pageNumber": None,
            "section": stem_name,
            "text": raw_text.strip(),
            "metadata": {"format": "md", "section_title": stem_name, "character_count": len(raw_text.strip())}
        })

    return sections, file_bytes


def extract_document_content(
    file_bytes: bytes,
    filename: str,
    mime_type: Optional[str] = None,
    password: Optional[str] = None
) -> Tuple[List[Dict[str, Any]], bytes]:
    """
    Universal entry point for multi-format extraction.
    Normalizes PDF, DOCX, TXT, and MD into structured sections with provenance.
    Returns: (sections_list, unlocked_or_clean_bytes)
    """
    fmt = sniff_document_type(file_bytes, filename, mime_type)

    if fmt == "pdf":
        return extract_pdf_content(file_bytes, filename, password=password)
    elif fmt == "docx":
        return extract_docx_content(file_bytes, filename)
    elif fmt == "txt":
        return extract_txt_content(file_bytes, filename)
    elif fmt == "md":
        return extract_markdown_content(file_bytes, filename)
    else:
        raise UnsupportedFormatError(f"Unsupported format: '{fmt}'")
